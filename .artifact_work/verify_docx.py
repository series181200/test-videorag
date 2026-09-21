from pathlib import Path
from zipfile import ZipFile
from docx import Document

path = Path(r"D:\file\file\华科\作业\软件质量测试\test-videorag\outputs\vimo_test_documents_20260915\附录7：Vimo Desktop测试脚本运行配置说明.docx")
doc = Document(path)
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text:
        print(f"P{i:03d} [{p.style.name}] {text}")
for ti, table in enumerate(doc.tables):
    print(f"TABLE {ti}: rows={len(table.rows)} cols={len(table.columns)}")
    for row in table.rows:
        print(" | ".join(cell.text.replace("\n", " / ") for cell in row.cells))

all_text = "\n".join(p.text for p in doc.paragraphs)
all_text += "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
required = [
    "测试脚本运行配置说明", "Vimo Desktop", "pytest", "Vitest",
    "run_all_layer_tests.py", "34", "32", "94.12%", "TC-COM-003", "TC-API-004"
]
missing = [item for item in required if item not in all_text]
print(f"missing_required={missing}")
with ZipFile(path) as archive:
    bad = archive.testzip()
    print(f"zip_test={bad}")
