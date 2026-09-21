from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"D:\file\file\华科\作业\软件质量测试\第一阶段\实践作业-文档模板2025\实践作业-文档模板2025\附录2：程序运行说明模板.docx")
OUTPUT = Path(r"D:\file\file\华科\作业\软件质量测试\test-videorag\outputs\vimo_test_documents_20260915\附录2：Vimo Desktop程序运行说明.docx")
EXPECTED_SHA256 = "83215ae3c5c6e409878b6697dea556bd25326de570706a0c5930923b18df1164"

BODY_STYLE = "石墨文档正文"
BLUE = RGBColor(0x00, 0x00, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = "FFFFFF"
NAVY = "1F4E78"
PALE_BLUE = "F2F6FB"
LIGHT_GRAY = "D9D9D9"


def set_font(run, name="微软雅黑", size=11, bold=None, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.35):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def add_body(doc, text="", bold=False, color=BLACK, before=0, after=6, align=None):
    p = doc.add_paragraph(style=BODY_STYLE)
    if align is not None:
        p.alignment = align
    set_paragraph_spacing(p, before=before, after=after)
    r = p.add_run(text)
    set_font(r, size=11, bold=bold, color=color)
    return p


def add_subheading(doc, text):
    p = doc.add_paragraph(style=BODY_STYLE)
    set_paragraph_spacing(p, before=8, after=4, line=1.15)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, size=12, bold=True, color=BLACK)
    return p


def add_step(doc, number, title, detail):
    p = doc.add_paragraph(style=BODY_STYLE)
    set_paragraph_spacing(p, before=2, after=4)
    r1 = p.add_run(f"{number}. {title}  ")
    set_font(r1, size=11, bold=True)
    r2 = p.add_run(detail)
    set_font(r2, size=11)
    return p


def add_code(doc, lines):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = p.paragraph_format
    fmt.left_indent = Inches(0.2)
    fmt.right_indent = Inches(0.1)
    fmt.space_before = Pt(2)
    fmt.space_after = Pt(7)
    fmt.line_spacing = 1.1
    fmt.keep_together = True
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F3F5F7")
    p_pr.append(shd)
    for index, line in enumerate(lines):
        run = p.add_run(line)
        set_font(run, name="Consolas", size=9.5, color=BLACK)
        if index < len(lines) - 1:
            run.add_break()
    return p


def add_table(doc, headers, rows, widths, add_spacer=True):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_repeat_header(table.rows[0])
    set_cant_split(table.rows[0])

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, after=0, line=1.1)
        r = p.add_run(header)
        set_font(r, size=10.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    for row_index, row_data in enumerate(rows):
        new_row = table.add_row()
        set_cant_split(new_row)
        cells = new_row.cells
        for col_index, value in enumerate(row_data):
            cell = cells[col_index]
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_spacing(p, after=0, line=1.2)
            r = p.add_run(str(value))
            set_font(r, size=10.5)

    if add_spacer:
        spacer = doc.add_paragraph(style=BODY_STYLE)
        set_paragraph_spacing(spacer, after=3, line=1.0)
    return table


def add_heading(doc, text, page_break=False):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break
    r = p.add_run(text)
    set_font(r, size=15, bold=True, color=BLACK)
    return p


def clear_document_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


actual_sha = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
if actual_sha != EXPECTED_SHA256:
    raise RuntimeError(f"模板已变化，期望 {EXPECTED_SHA256}，实际 {actual_sha}")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
shutil.copy2(SOURCE, OUTPUT)
doc = Document(OUTPUT)
clear_document_body(doc)

# 标题和基本信息
title = doc.add_paragraph(style=BODY_STYLE)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(12)
run = title.add_run("程序运行说明")
set_font(run, size=16, bold=True, color=BLACK)

for label, value, value_color in (
    ("班级：", "待填写", BLUE),
    ("产品主题：", "Vimo Desktop 视频理解与问答应用", BLUE),
    ("小组长：", "待填写", BLUE),
):
    p = doc.add_paragraph(style=BODY_STYLE)
    set_paragraph_spacing(p, after=3, line=1.15)
    r1 = p.add_run(label)
    set_font(r1, size=11)
    r2 = p.add_run(value)
    set_font(r2, size=11, color=value_color)

intro = add_body(
    doc,
    "本说明用于在 Windows 环境中从源代码启动 Vimo Desktop。程序由 Electron 桌面端和 Python VideoRAG 服务组成，开发模式下必须先启动后端，再启动桌面端。",
    before=5,
    after=8,
)

# 第一部分
add_heading(doc, "1 程序结构说明")
add_body(
    doc,
    "Vimo Desktop 按调用链分为页面层、桌面通信层、Python API 层和 VideoRAG 算法层。用户在页面中导入视频和提问，Electron 主进程通过 IPC 与本地 HTTP 接口转发请求，Python 后端负责异步任务与状态管理，算法层完成视频分段、特征提取、索引和检索问答。",
)

add_table(
    doc,
    ["目录或文件", "作用"],
    [
        ("src/renderer", "React 页面、初始化向导、视频导入、会话和问答界面"),
        ("src/preload 与 src/main", "通过 contextBridge 暴露受控 API，并由 Electron 主进程完成 IPC、配置读写和 Python API 转发"),
        ("python_backend", "videorag_api.py 提供 Flask API 和任务管理；videorag 子目录实现视频处理、多模态索引和检索生成算法"),
    ],
    [2.15, 3.55],
    add_spacer=False,
)

# 第二部分
add_heading(doc, "2 程序运行方法", page_break=True)
add_subheading(doc, "2.1 运行环境")
add_table(
    doc,
    ["项目", "建议版本或要求", "说明"],
    [
        ("操作系统", "Windows 10 或 11 64 位", "本说明采用 PowerShell；macOS 使用相同命令但路径分隔符为 /"),
        ("Python", "3.11", "建议使用 Conda 创建名为 vimo 的独立环境"),
        ("Node.js", ">= 20", "package.json 的最低版本要求"),
        ("pnpm", "9.10.0", "项目锁文件和 packageManager 指定的包管理器"),
        ("Git", "可用版本", "安装 pytorchvideo 和 ImageBind 源码依赖时使用"),
        ("网络与磁盘", "稳定网络，建议至少 10 GB 空闲空间", "ImageBind 模型约 4.5 GB，依赖和分析缓存还会占用空间"),
        ("FFmpeg", "建议安装并加入 PATH", "MoviePy 处理视频和音频时需要"),
    ],
    [1.15, 1.55, 3.0],
)
add_body(doc, "CPU 可以运行程序，但视频分析和 ImageBind 加载会较慢；如使用 NVIDIA GPU，应安装与本机 CUDA 环境匹配的 PyTorch。", after=5)

add_subheading(doc, "2.2 创建 Python 环境")
p = add_body(doc, "在 PowerShell 中执行以下命令。首次安装需要联网，Git 依赖下载时间较长。")
p.paragraph_format.keep_with_next = True
add_code(doc, [
    "conda create --name vimo python=3.11 -y",
    "conda activate vimo",
    "python -m pip install --upgrade pip",
])
add_code(doc, [
    "python -m pip install numpy==1.26.4 torch==2.1.2 torchvision==0.16.2 torchaudio==2.1.2",
    "python -m pip install moviepy==1.0.3",
    "python -m pip install git+https://github.com/Re-bin/pytorchvideo.git@58f50da4e4b7bf0b17b1211dc6b283ba42e522df",
    "python -m pip install --no-deps git+https://github.com/facebookresearch/ImageBind.git@3fcf5c9039de97f6ff5528ee4a9dce903c5979b3",
])
add_code(doc, [
    "python -m pip install timm ftfy regex einops fvcore eva-decord==0.6.1 iopath matplotlib types-regex cartopy",
    "python -m pip install neo4j hnswlib xxhash nano-vectordb",
    "python -m pip install tiktoken openai tenacity dashscope",
    "python -m pip install flask psutil flask_cors setproctitle pywin32",
])

add_subheading(doc, "2.3 安装桌面端依赖")
p = add_body(doc, "进入 Vimo-desktop 根目录，使用项目指定的 pnpm 安装依赖。")
p.paragraph_format.keep_with_next = True
add_code(doc, [
    "cd <项目目录>\\Vimo-desktop",
    "corepack enable",
    "corepack prepare pnpm@9.10.0 --activate",
    "pnpm install",
])
add_body(doc, "若当前 Node.js 未提供 Corepack，可执行 npm.cmd install -g pnpm@9.10.0 后再运行 pnpm install。", after=5)

add_subheading(doc, "2.4 启动 Python 后端")
p = add_body(doc, "打开第一个 PowerShell 窗口并保持运行：")
p.paragraph_format.keep_with_next = True
add_code(doc, [
    "conda activate vimo",
    "cd <项目目录>\\Vimo-desktop\\python_backend",
    "python videorag_api.py",
])
add_body(
    doc,
    "服务默认使用 64451 端口；该端口被占用时，会在 64451 至 64470 范围内选择可用端口。出现 VideoRAG API is running 或 Flask listening 信息后再启动桌面端。",
)

add_subheading(doc, "2.5 启动桌面端")
p = add_body(doc, "打开第二个 PowerShell 窗口：")
p.paragraph_format.keep_with_next = True
add_code(doc, [
    "cd <项目目录>\\Vimo-desktop",
    "pnpm dev",
])
add_body(doc, "开发模式下 Electron 会扫描已启动的后端服务。若后端未运行，界面会持续等待或提示未找到服务。", after=5)

add_subheading(doc, "2.6 首次启动配置")
add_step(doc, 1, "选择存储目录", "选择有写权限且空间充足的目录，用于保存模型、配置、会话和索引数据。")
add_step(doc, 2, "准备 ImageBind", "点击 Start Download 下载约 4.5 GB 的模型，等待状态变为 Completed。模型文件保存在 <存储目录>\\imagebind_huge\\imagebind_huge.pth。")
add_step(doc, 3, "配置 OpenAI", "填写 Base URL、OpenAI API Key、Processing Model 和 Analysis Model。默认模型名为 gpt-4o-mini。")
add_step(doc, 4, "配置 DashScope", "填写 DashScope API Key。默认视频描述模型为 qwen-vl-plus-latest，语音识别模型为 paraformer-realtime-v2。")
add_step(doc, 5, "完成初始化", "保存配置后进入主界面。当前后端实际运行视频分析时要求 OpenAI 和 DashScope 两项 API Key 均已配置。")
add_body(
    doc,
    "程序没有内置登录账户或密码。API Key 必须由使用者从对应服务商获取，文档和截图中不要记录真实密钥。存储目录记录在 %USERPROFILE%\\.videorag-bootstrap.json，其他设置保存在所选目录的 config.json。",
    after=6,
)

add_subheading(doc, "2.7 主要操作流程")
add_step(doc, 1, "新建会话", "在侧边栏创建或选择一个会话。")
add_step(doc, 2, "导入视频", "选择一个或多个受支持的视频文件，确认文件名和路径正确。")
add_step(doc, 3, "开始分析", "提交视频后等待索引状态从 processing 变为 completed。处理时间取决于视频长度、硬件和模型服务速度。")
add_step(doc, 4, "视频问答", "分析完成后输入问题，等待回答显示在当前会话中。")
add_step(doc, 5, "保存与退出", "会话会保存到存储目录。退出后不要移动或重命名已导入的视频，否则再次访问时可能找不到源文件。")

add_subheading(doc, "2.8 停止程序与常见问题")
add_body(doc, "先关闭 Electron 窗口，再分别在两个 PowerShell 窗口按 Ctrl+C 结束桌面开发进程和 Python 后端。")
doc.add_page_break()
add_table(
    doc,
    ["现象", "检查方法"],
    [
        ("桌面端一直等待后端", "确认 python videorag_api.py 正在运行，并检查 64451 至 64470 端口是否被防火墙阻止"),
        ("提示缺少配置", "确认存储目录、OpenAI API Key、DashScope API Key 均已保存"),
        ("ImageBind 加载失败", "确认 imagebind_huge.pth 位于存储目录下的 imagebind_huge 文件夹，且磁盘空间和文件权限正常"),
        ("视频处理失败", "确认 FFmpeg 可用、视频路径未变化，并查看 Python 后端终端中的错误信息"),
        ("pnpm 命令不可用", "确认 Node.js >= 20，并通过 Corepack 或 npm 全局安装 pnpm 9.10.0"),
    ],
    [2.0, 3.7],
)
add_body(doc, "macOS 的启动顺序相同，路径改用 /。PyTorch 和系统依赖应按 Mac 硬件重新选择，不能直接照搬 Windows GPU 安装方式。", after=4)

# 第三部分
add_heading(doc, "3 程序运行效果", page_break=True)
add_body(doc, "程序运行成功后，可通过以下现象确认完整链路可用：")
add_table(
    doc,
    ["阶段", "预期效果"],
    [
        ("后端启动", "Python 终端显示 Flask 服务端口，访问 /api/health 可得到 status=ok"),
        ("桌面端启动", "弹出 Vimo Desktop 窗口；首次运行显示初始化向导，完成后进入主界面"),
        ("模型准备", "ImageBind 下载进度达到 100%，服务状态显示可进行视频处理"),
        ("视频导入", "界面显示所选视频的文件名、路径和大小，重复选择不会产生重复项"),
        ("视频分析", "界面显示 processing 状态和当前步骤，完成后显示已索引视频"),
        ("视频问答", "用户问题出现在会话中，查询完成后显示模型生成的回答"),
        ("会话恢复", "关闭并重新启动应用后，已保存的会话标题、消息和视频信息仍可读取"),
    ],
    [1.45, 4.25],
)
# 页面属性沿用模板，只统一标题和正文样式。
for section in doc.sections:
    section.header.is_linked_to_previous = True
    section.footer.is_linked_to_previous = True

doc.core_properties.title = "Vimo Desktop 程序运行说明"
doc.core_properties.subject = "软件质量测试实践课程程序运行说明"
doc.core_properties.keywords = "Vimo Desktop, VideoRAG, Electron, Python"
doc.save(OUTPUT)
print(OUTPUT)
