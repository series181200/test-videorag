from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

SOURCE = Path(r"D:\file\file\华科\作业\软件质量测试\第一阶段\实践作业-文档模板2025\实践作业-文档模板2025\附录7：测试脚本运行配置说明模板.docx")
OUTPUT_DIR = Path(r"D:\file\file\华科\作业\软件质量测试\test-videorag\outputs\vimo_test_documents_20260915")
OUTPUT = OUTPUT_DIR / "附录7：Vimo Desktop测试脚本运行配置说明.docx"
RESULT_IMAGE = OUTPUT_DIR.parent.parent / ".artifact_work" / "四层测试结果摘要.png"


def set_run_font(run, name="宋体", size=10.5, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_paragraph(p, before=0, after=6, line=1.25):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    for run in p.runs:
        set_run_font(run)


def add_text(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        a = p.add_run(bold_prefix)
        set_run_font(a, bold=True)
        b = p.add_run(text[len(bold_prefix):])
        set_run_font(b)
    else:
        set_run_font(p.add_run(text))
    style_paragraph(p)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.55)
    p.paragraph_format.first_line_indent = Cm(-0.3)
    set_run_font(p.add_run("• "))
    set_run_font(p.add_run(text))
    style_paragraph(p, after=3)
    return p


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, "F2F2F2")
    set_cell_margins(cell, 120, 140, 120, 140)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for index, line in enumerate(lines):
        if index:
            p.add_run().add_break()
        set_run_font(p.add_run(line), name="Consolas", size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = text
        shade_cell(hdr[i], "1F4E78")
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            set_run_font(run, size=9.5, bold=True, color=(255, 255, 255))
        set_cell_margins(hdr[i])
    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.08
                for run in p.runs:
                    set_run_font(run, size=9)
            set_cell_margins(cells[i])
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc, text, level=1):
    if level == 1:
        p = doc.add_heading(text, level=1)
    else:
        p = doc.add_paragraph()
        p.add_run(text)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(4)
    for run in p.runs:
        set_run_font(run, name="黑体", size=14 if level == 1 else 12, bold=True)
    return p


def build_result_image():
    width, height = 1500, 560
    image = Image.new("RGB", (width, height), "#111827")
    draw = ImageDraw.Draw(image)
    font_path = r"C:\Windows\Fonts\msyh.ttc"
    font = ImageFont.truetype(font_path, 30)
    small = ImageFont.truetype(font_path, 25)
    green = "#86EFAC"
    red = "#FCA5A5"
    white = "#F9FAFB"
    gray = "#CBD5E1"
    lines = [
        ("Vimo 四层统一测试结果", white, font),
        ("Renderer 页面层          总数  5    OK  5    NG  0", green, small),
        ("前后端通信层             总数  8    OK  7    NG  1", red, small),
        ("后端与算法通信层         总数 14    OK 13    NG  1", red, small),
        ("VideoRAG 算法层          总数  7    OK  7    NG  0", green, small),
        ("总计                     总数 34    OK 32    NG  2    通过率 94.12%", white, small),
        ("未通过：TC-COM-003 Windows 文件名提取；TC-API-004 非字符串 query 返回 500", gray, small),
    ]
    y = 42
    for text, color, current_font in lines:
        draw.text((55, y), text, font=current_font, fill=color)
        y += 68
    RESULT_IMAGE.parent.mkdir(parents=True, exist_ok=True)
    image.save(RESULT_IMAGE)


doc = Document(SOURCE)
body = doc._element.body
sect_pr = body.sectPr
for child in list(body):
    if child is not sect_pr:
        body.remove(child)

section = doc.sections[0]
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.3)
section.right_margin = Cm(2.3)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(14)
set_run_font(title.add_run("测试脚本运行配置说明"), name="黑体", size=18, bold=True)

meta = [
    ("班级：", "待填写"),
    ("产品主题：", "Vimo Desktop 视频检索增强问答应用"),
    ("小组长：", "待填写"),
]
for label, value in meta:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run(label), bold=True)
    value_run = p.add_run(value)
    set_run_font(value_run, color=(0, 0, 255) if value == "待填写" else None)

add_text(doc, "本说明用于配置和运行 Vimo Desktop 四层自动化测试。统一入口依次执行 Renderer 页面层、Electron 通信层、Python API 层和 VideoRAG 算法层，并逐条输出测试层级、接口或函数、OK/NG 状态和缺陷说明。GUI 测试与功能测试在下文分别说明，但由同一入口统一调度。")

add_heading(doc, "1 测试工具概述")
add_table(
    doc,
    ["工具", "版本", "用途", "获取方式"],
    [
        ["Python", "3.11.16", "运行统一脚本及 Python 测试", "Conda 环境 vimo"],
        ["pytest", "8.4.2", "API 层和算法层单元测试", "https://pytest.org"],
        ["Node.js", "24.21.0", "运行 TypeScript 测试", "https://nodejs.org"],
        ["npm", "11.19.0", "安装前端与通信层依赖", "随 Node.js 安装"],
        ["Vitest", "2.1.9", "Renderer 层和通信层单元测试", "https://vitest.dev"],
        ["Testing Library", "React 16.3.0", "模拟用户点击、输入和页面断言", "test_renderer 依赖"],
        ["jsdom", "26.1.0", "在 Node.js 中模拟浏览器 DOM", "test_renderer 依赖"],
    ],
    [2.5, 2.2, 5.6, 5.2],
)
add_text(doc, "TypeScript 测试统一使用 Vitest，Python 测试统一使用 pytest。Renderer 测试使用 jsdom 和 Testing Library，不需要安装 Playwright 或 Chromium。测试通过 mock 隔离 Electron、文件系统、HTTP 请求、模型服务和子进程，因此日常单元测试不需要真实 OpenAI 或 DashScope API Key。")

add_heading(doc, "2 测试脚本说明")
add_heading(doc, "2.1 运行环境", level=2)
add_bullet(doc, "操作系统：Windows 10 或 Windows 11，使用 PowerShell 或命令提示符。")
add_bullet(doc, "工作目录：Vimo-desktop 项目根目录。所有命令均从该目录执行。")
add_bullet(doc, "Python 环境：建议激活 vimo Conda 环境；本次有效执行使用 Python 3.11.16。")
add_bullet(doc, "Node.js 环境：本次执行使用 Node.js 24.21.0 和 npm 11.19.0。")
add_bullet(doc, "账户与密码：单元测试不需要真实账户、密码或 API Key。测试中的配置和密钥均为 mock 数据。")

add_heading(doc, "2.2 首次安装", level=2)
add_text(doc, "在 Vimo-desktop 根目录激活 vimo 环境后执行：")
add_code_block(doc, [
    "python -m pip install -r requirements-unified-tests.txt",
    "npm.cmd install --prefix test_renderer",
    "npm.cmd install --prefix test_communication",
])
add_text(doc, "如果 python 指向 WindowsApps 占位程序或其他未安装 pytest 的解释器，应先激活 vimo 环境，或使用 VIMO_TEST_PYTHON 指定解释器。")
add_code_block(doc, [
    '$env:VIMO_TEST_PYTHON = "C:\\Users\\24032\\.conda\\envs\\vimo\\python.exe"',
    ".\\run_all_layer_tests.bat",
])

add_heading(doc, "2.3 一键运行", level=2)
add_text(doc, "推荐双击 run_all_layer_tests.bat。也可以在已经激活 vimo 环境的终端中执行：")
add_code_block(doc, ["python run_all_layer_tests.py"])
add_text(doc, "统一脚本在某一层失败后仍会继续执行后续层。最终按层汇总用例总数、OK、NG，并以退出码 0 表示全部通过，以非 0 表示存在失败项或测试环境问题。")

add_heading(doc, "2.4 分层脚本和测试目的", level=2)
add_table(
    doc,
    ["测试层级", "脚本范围", "框架", "用例数", "测试目的"],
    [
        ["Renderer 页面层", "test_renderer", "Vitest", "5", "验证初始化、视频选择、分析、问答和错误反馈等用户可见流程"],
        ["前后端通信层", "test_communication/filtered", "Vitest", "8", "验证 preload API、IPC 映射、文件和配置通信、会话保存及后端错误转换"],
        ["后端与算法通信层", "test_videorag_api/tests", "pytest", "14", "验证 Flask 接口、Worker 异常、进程生命周期、日志安全和状态持久化"],
        ["VideoRAG 算法层", "test_videorag_algorithm/filtered", "pytest", "7", "验证查询默认值、JSON 提取、Token 边界、分块、持久化和并发恢复"],
    ],
    [2.7, 3.7, 1.8, 1.3, 6.2],
)

add_heading(doc, "2.5 分层单独运行", level=2)
add_code_block(doc, [
    "test_renderer\\node_modules\\.bin\\vitest.cmd run --config test_renderer\\vitest.config.ts",
    "test_communication\\node_modules\\.bin\\vitest.cmd run --config test_communication\\filtered\\vitest.config.ts",
    "python -m pytest -c test_videorag_api\\pytest.ini test_videorag_api\\tests",
    "python -m pytest -c test_videorag_algorithm\\filtered\\pytest.ini test_videorag_algorithm\\filtered",
])

add_heading(doc, "2.6 结果判定", level=2)
add_bullet(doc, "OK：实际结果满足断言，终端显示未发现漏洞。")
add_bullet(doc, "NG：实际结果与预期不一致，终端显示对应模块和缺陷说明。")
add_bullet(doc, "环境问题：依赖缺失、Python 解释器错误或临时目录不可写。该结果应与业务缺陷分开记录。")

add_heading(doc, "3 测试脚本运行效果")
add_text(doc, "在依赖完整的 vimo 环境中执行当前代码，34 条用例中 32 条通过、2 条失败，通过率为 94.12%。各层结果如下。")
add_table(
    doc,
    ["测试层级", "总数", "OK", "NG", "通过率"],
    [
        ["Renderer 页面层", "5", "5", "0", "100.00%"],
        ["前后端通信层", "8", "7", "1", "87.50%"],
        ["后端与算法通信层", "14", "13", "1", "92.86%"],
        ["VideoRAG 算法层", "7", "7", "0", "100.00%"],
        ["合计", "34", "32", "2", "94.12%"],
    ],
    [5.0, 2.0, 2.0, 2.0, 3.0],
)
build_result_image()
figure = doc.add_paragraph()
figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
figure.add_run().add_picture(str(RESULT_IMAGE), width=Cm(15.5))
caption = doc.add_paragraph("图1 四层统一测试结果摘要")
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
caption.paragraph_format.space_after = Pt(6)
for run in caption.runs:
    set_run_font(run, size=9)

add_heading(doc, "3.1 当前未通过项", level=2)
add_table(
    doc,
    ["用例", "模块", "实际结果", "影响"],
    [
        ["TC-COM-003", "select-video-files", "Windows 路径返回完整路径作为 name", "页面和会话标题可能显示本地完整路径"],
        ["TC-API-004", "POST /api/sessions/<chat_id>/query", "非字符串 query 返回 HTTP 500，预期为 400", "无效输入进入内部处理，错误分类不正确"],
    ],
    [2.0, 3.4, 5.4, 4.8],
)

add_heading(doc, "3.2 缺陷修复回归", level=2)
add_text(doc, "已对 TC-COM-002、TC-API-009、TC-API-011 和 TC-ALG-007 实施修复。单独回归脚本 run_fixed_defect_tests.py 的结果为 4 条全部通过，Electron 生产构建同时通过。")
add_table(
    doc,
    ["用例", "修复内容", "回归结果"],
    [
        ["TC-COM-002", "补充 preload 通用 invoke，并设置 IPC 通道白名单", "OK"],
        ["TC-API-009", "终止会话时同时清理索引和查询进程", "OK"],
        ["TC-API-011", "日志递归脱敏 API Key、token、secret 和 password 字段", "OK"],
        ["TC-ALG-007", "使用 try/finally 保证异常和取消后释放并发容量", "OK"],
    ],
    [2.5, 11.0, 2.0],
)
add_text(doc, "当前核心回归集能够快速验证四层主要业务链。真实模型调用、GPU 资源、超大视频、网络中断和长时间高并发不属于本次单元测试范围，应在集成测试、性能测试和端到端测试中补充。")

for paragraph in doc.paragraphs:
    if paragraph.style.name != "Heading 1" and not (
        paragraph.runs and paragraph.runs[0].bold and paragraph.runs[0].font.name == "黑体"
    ):
        style_paragraph(paragraph, after=paragraph.paragraph_format.space_after.pt if paragraph.paragraph_format.space_after else 5)

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
