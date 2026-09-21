from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

source = Path(r"D:\file\file\华科\作业\软件质量测试\第一阶段\实践作业-文档模板2025\实践作业-文档模板2025\附录7：测试脚本运行配置说明模板.docx")
doc = Document(source)

print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")
for i, p in enumerate(doc.paragraphs):
    parts = []
    for run in p.runs:
        color = run.font.color.rgb
        parts.append(f"[{run.text!r}|color={color}|bold={run.bold}|size={run.font.size}]")
    print(f"P{i:03d} style={p.style.name!r}: {''.join(parts)}")

for ti, table in enumerate(doc.tables):
    print(f"TABLE {ti} rows={len(table.rows)} cols={len(table.columns)}")
    for ri, row in enumerate(table.rows):
        cells = []
        for ci, cell in enumerate(row.cells):
            text = "\\n".join(p.text for p in cell.paragraphs)
            colors = []
            for p in cell.paragraphs:
                for run in p.runs:
                    colors.append(str(run.font.color.rgb))
            cells.append(f"C{ci}={text!r} colors={colors}")
        print(f"R{ri}: " + " || ".join(cells))

for si, section in enumerate(doc.sections):
    print(f"SECTION {si}: size={section.page_width}x{section.page_height} margins={section.top_margin},{section.right_margin},{section.bottom_margin},{section.left_margin}")
